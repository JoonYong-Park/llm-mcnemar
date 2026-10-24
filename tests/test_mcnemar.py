import json
from types import SimpleNamespace

import pandas as pd
import pytest
from openpyxl import load_workbook

import mcnemar

TS = "2025-01-01T00-00-00.000000"


def _em_frame(a, b):
    return pd.DataFrame({"m_a_EM": a, "m_b_EM": b})


def test_mcnemar_significant():
    # only A correct: 40, only B correct: 5 -> significant
    a = [1] * 40 + [0] * 5 + [1] * 50
    b = [0] * 40 + [1] * 5 + [1] * 50
    row = mcnemar.mcnemar_test(_em_frame(a, b), "m_a", "m_b")
    assert row["Questions"] == 95
    assert row["Only m_a correct"] == 40
    assert row["Only m_b correct"] == 5
    assert row["Both correct"] == 50
    assert row["p-value"] < 0.05
    assert row["Significance"] == mcnemar.SIGNIFICANT


def test_mcnemar_not_significant():
    a = [1] * 20 + [0] * 18
    b = [0] * 20 + [1] * 18
    assert mcnemar.mcnemar_test(_em_frame(a, b), "m_a", "m_b")["Significance"] == mcnemar.NOT_SIGNIFICANT


def test_mcnemar_too_few_discordant_pairs():
    a = [1] * 10 + [0] * 2
    b = [0] * 10 + [1] * 2
    assert mcnemar.mcnemar_test(_em_frame(a, b), "m_a", "m_b")["Significance"].startswith("N/A")
    row = mcnemar.mcnemar_test(_em_frame(a, b), "m_a", "m_b", min_discordant=5)
    assert not row["Significance"].startswith("N/A")


def _write_samples(path, task, scores, ts=TS):
    path.mkdir(parents=True, exist_ok=True)
    with open(path / f"samples_cot_{task}_{ts}.jsonl", "w", encoding="utf-8") as f:
        for s in scores:
            f.write(json.dumps({"exact_match": s}) + "\n")


def _config(data_path, **overrides):
    values = {
        **mcnemar._CONFIG_DEFAULTS,
        "NAME": "test",
        "DATA_DIR": "test",
        "RESULT_SLUG": "test",
        "BENCHMARKS": ["bench1", "bench2"],
        "CATEGORY_BENCHMARKS": ["bench1"],
        "DOMAIN_TO_BENCHMARKS": {"Domain 1": ["bench1_x", "bench2"]},
        "DOMAINS": ["Domain 1"],
        **overrides,
    }
    cfg = SimpleNamespace(**values)
    cfg.DATA_PATH = data_path
    return cfg


@pytest.fixture
def data_path(tmp_path):
    root = tmp_path / "data"
    # bench1: 2 categories (x, y); x also has an older run that must be ignored
    _write_samples(root / "model_a" / "bench1", "bench1_x", [1] * 40 + [0] * 10)
    _write_samples(root / "model_a" / "bench1", "bench1_x", [0] * 50, ts="2024-01-01T00-00-00.000000")
    _write_samples(root / "model_a" / "bench1", "bench1_y", [1] * 30)
    _write_samples(root / "model_b" / "bench1", "bench1_x", [0] * 40 + [0] * 10)
    _write_samples(root / "model_b" / "bench1", "bench1_y", [1] * 30)
    # bench2: question counts differ -> skipped
    _write_samples(root / "model_a" / "bench2", "bench2", [1] * 10)
    _write_samples(root / "model_b" / "bench2", "bench2", [1] * 12)
    return root


def test_run_end_to_end(data_path, tmp_path):
    out = mcnemar.run(_config(data_path), "model_a", "model_b", output_dir=tmp_path / "result")

    wb = load_workbook(out)
    assert wb.sheetnames == ["accuracy", "benchmark_mcnemar", "domain_mcnemar", "bench1", "raw", "info"]
    assert wb["accuracy"].freeze_panes == "B2"

    bench = pd.read_excel(out, sheet_name="benchmark_mcnemar")
    assert list(bench["Benchmark"]) == ["bench1"]  # bench2 skipped (question count mismatch)
    assert bench.loc[0, "Only model_a correct"] == 40  # latest timestamp file is used
    assert bench.loc[0, "Significance"] == mcnemar.SIGNIFICANT

    cats = pd.read_excel(out, sheet_name="bench1").set_index("Category")
    assert set(cats.index) == {"bench1_x", "bench1_y"}
    assert cats.loc["bench1_y", "Significance"].startswith("N/A")

    acc = pd.read_excel(out, sheet_name="accuracy").set_index("Name")
    assert list(acc["Type"]) == ["Benchmark", "Domain", "Category", "Category"]
    assert acc.loc["bench1", "model_a accuracy (%)"] == pytest.approx(87.5)
    assert acc.loc["bench1", "Diff (A-B, %p)"] == pytest.approx(50.0)
    assert acc.loc["Domain 1", "Questions"] == 50

    info = pd.read_excel(out, sheet_name="info").set_index("Item")["Value"]
    assert info["Benchmarks skipped"] == 1


def test_extra_groups_and_labels(data_path, tmp_path):
    cfg = _config(
        data_path,
        BENCHMARK_LABELS={"bench1": "B1"},
        EXTRA_GROUPS={"bench1_x_only": {"benchmarks": ["bench1"], "category_contains": "_x"}},
    )
    out = mcnemar.run(cfg, "model_a", "model_b", output_dir=tmp_path / "result")
    bench = pd.read_excel(out, sheet_name="benchmark_mcnemar").set_index("Benchmark")
    assert list(bench.index) == ["B1", "bench1_x_only"]
    assert bench.loc["bench1_x_only", "Both correct"] == 0


def test_available_configs():
    assert {"legal", "korean", "english", "example"} <= set(mcnemar.available_configs())
    for name in mcnemar.available_configs():
        mcnemar.load_config(name)


def test_example_data(tmp_path):
    """Full run on data/example + configs/example.py."""
    out = mcnemar.run(mcnemar.load_config("example"), "model-a", "model-b", output_dir=tmp_path)

    bench = pd.read_excel(out, sheet_name="benchmark_mcnemar").set_index("Benchmark")
    assert "winogrande" not in bench.index  # question count mismatch -> skipped
    assert bench.loc["kmmlu", "Significance"] == mcnemar.SIGNIFICANT
    assert bench.loc["gsm8k", "Questions"] == 300  # only one filter is kept
    assert {"gsm8k+hellaswag", "kmmlu_math_only"} <= set(bench.index)


def _write_rows(path, task, rows):
    path.mkdir(parents=True, exist_ok=True)
    with open(path / f"samples_{task}_{TS}.jsonl", "w", encoding="utf-8") as f:
        for r in rows:
            f.write(json.dumps(r) + "\n")


def test_pairs_by_doc_id_when_order_differs(tmp_path):
    # model_b lists the same questions in reverse order; pairing must follow doc_id, not row order
    scores_a = [1] * 30 + [0] * 30
    scores_b = [0] * 30 + [1] * 30
    rows_a = [{"doc_id": i, "exact_match": s} for i, s in enumerate(scores_a)]
    rows_b = [{"doc_id": i, "exact_match": s} for i, s in enumerate(scores_b)][::-1]
    _write_rows(tmp_path / "data" / "model_a" / "bench1", "bench1", rows_a)
    _write_rows(tmp_path / "data" / "model_b" / "bench1", "bench1", rows_b)

    cfg = _config(tmp_path / "data", BENCHMARKS=["bench1"], CATEGORY_BENCHMARKS=[],
                  DOMAIN_TO_BENCHMARKS={}, DOMAINS=[])
    out = mcnemar.run(cfg, "model_a", "model_b", output_dir=tmp_path / "result")
    bench = pd.read_excel(out, sheet_name="benchmark_mcnemar").set_index("Benchmark")
    assert bench.loc["bench1", "Only model_a correct"] == 30
    assert bench.loc["bench1", "Only model_b correct"] == 30
    assert bench.loc["bench1", "Both correct"] == 0


def test_non_binary_scores_are_skipped(tmp_path):
    _write_rows(tmp_path / "data" / "model_a" / "bench1", "bench1",
                [{"doc_id": i, "exact_match": 0.5} for i in range(10)])
    _write_rows(tmp_path / "data" / "model_b" / "bench1", "bench1",
                [{"doc_id": i, "exact_match": 1} for i in range(10)])
    cfg = _config(tmp_path / "data", BENCHMARKS=["bench1"], CATEGORY_BENCHMARKS=[],
                  DOMAIN_TO_BENCHMARKS={}, DOMAINS=[])
    assert mcnemar.run(cfg, "model_a", "model_b", output_dir=tmp_path / "result") is None
