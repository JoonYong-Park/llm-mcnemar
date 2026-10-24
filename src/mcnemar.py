"""McNemar's test for LLM benchmark results.

Pairs two models' per-question correct/incorrect results, runs McNemar's test at the
benchmark, category and domain level, and writes an Excel report to
result/{A}__vs__{B}__{RESULT_SLUG}.xlsx.

Usage:
    python src/mcnemar.py --config example -a model-a -b model-b
"""
import argparse
import importlib
import sys
from datetime import datetime
from pathlib import Path
from types import SimpleNamespace

import numpy as np
import pandas as pd
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter
from statsmodels.stats.contingency_tables import mcnemar

from loaders import load_model_result

REPO_ROOT = Path(__file__).resolve().parent.parent
CONFIG_DIR = REPO_ROOT / "src" / "configs"
DATA_ROOT = REPO_ROOT / "data"
RESULT_DIR = REPO_ROOT / "result"

_CONFIG_REQUIRED = [
    "DATA_DIR", "RESULT_SLUG", "BENCHMARKS", "DOMAINS",
    "DOMAIN_TO_BENCHMARKS", "CATEGORY_BENCHMARKS",
]
_CONFIG_DEFAULTS = {
    "DOMAIN_AVERAGE": {},
    "BENCH_FOLDERS": {},
    "SINGLE_FILE_BENCH": set(),
    "BENCHMARK_LABELS": {},
    "EXTRA_GROUPS": {},
    "ALPHA": 0.05,
    "MIN_DISCORDANT": 25,
}

# ifeval-style benchmarks: the benchmark score is the simple mean of its
# 4 sub-categories (plsa / ilsa / plla / illa)
IFEVAL_LIKE = ["ko_ifeval", "ifeval", "ifbench", "koifeval"]

# ── Report layout ──
SHEET_ACCURACY = "accuracy"
SHEET_BENCHMARK = "benchmark_mcnemar"
SHEET_DOMAIN = "domain_mcnemar"
SHEET_RAW = "raw"
SHEET_INFO = "info"

SIGNIFICANT = "Significant"
NOT_SIGNIFICANT = "Not significant"
SIG_COL = "Significance"
DIFF_COL = "Diff (A-B, %p)"
TYPE_ORDER = ["Benchmark", "Domain", "Category"]

HEADER_FILL = PatternFill(start_color="D9D9D9", end_color="D9D9D9", fill_type="solid")
ACCURACY_FILL = PatternFill(start_color="CCECFF", end_color="CCECFF", fill_type="solid")
SIGNIFICANT_FILL = PatternFill(start_color="FF9999", end_color="FF9999", fill_type="solid")
# p-values below 0.0001 display as "<0.0001" (the cell keeps the exact value)
NUMBER_FORMATS = {"Statistic": "0.00", "p-value": '[<0.0001]"<0.0001";0.0000', DIFF_COL: "0.00"}


def correct_col(model):
    return f"{model} correct"


def accuracy_col(model):
    return f"{model} accuracy (%)"


# ── Config ──
def available_configs():
    return sorted(p.stem for p in CONFIG_DIR.glob("*.py") if not p.stem.startswith("_"))


def load_config(name, data_dir=None):
    """configs/{name}.py -> config object, with DATA_PATH (folder holding the model folders) filled in."""
    module = importlib.import_module(f"configs.{name}")
    missing = [k for k in _CONFIG_REQUIRED if not hasattr(module, k)]
    if missing:
        raise SystemExit(f"[ERROR] configs/{name}.py is missing required keys: {missing}")
    values = {k: v for k, v in vars(module).items() if k.isupper()}
    cfg = SimpleNamespace(**{**_CONFIG_DEFAULTS, **values})
    cfg.NAME = name
    cfg.DATA_PATH = Path(data_dir).resolve() if data_dir else DATA_ROOT / cfg.DATA_DIR
    return cfg


def build_result_filename(model_A: str, model_B: str, slug: str) -> str:
    return f"{model_A}__vs__{model_B}__{slug}.xlsx"


# ── Group selection ──
def select_group(final_df: pd.DataFrame, key: str):
    """Rows of benchmark `key`, or of category `key`. None if neither exists."""
    if key in final_df["benchmark"].unique():
        return final_df[final_df["benchmark"] == key]
    if key in final_df["category"].unique():
        return final_df[final_df["category"] == key]
    return None


def extra_group(final_df: pd.DataFrame, spec: dict) -> pd.DataFrame:
    """EXTRA_GROUPS entry -> rows of the listed benchmarks, optionally filtered by a category token."""
    sub = final_df[final_df["benchmark"].isin(spec["benchmarks"])]
    token = spec.get("category_contains")
    if token:
        sub = sub[sub["category"].astype(str).str.contains(token, regex=False)]
    return sub


def domain_group(cfg, final_df: pd.DataFrame, domain_name: str) -> pd.DataFrame:
    """Union of the domain's benchmarks/categories (deduplicated by id)."""
    groups = [
        g for g in (select_group(final_df, key) for key in cfg.DOMAIN_TO_BENCHMARKS.get(domain_name, []))
        if g is not None
    ]
    if not groups:
        return final_df.iloc[0:0]
    return pd.concat(groups, ignore_index=True).drop_duplicates(subset="id")


# ── McNemar's test ──
def mcnemar_test(group: pd.DataFrame, model_A: str, model_B: str,
                 alpha: float = 0.05, min_discordant: int = 25):
    a_em, b_em = group[f"{model_A}_EM"], group[f"{model_B}_EM"]
    a_only = ((a_em == 1) & (b_em == 0)).sum()
    b_only = ((a_em == 0) & (b_em == 1)).sum()
    both_correct = ((a_em == 1) & (b_em == 1)).sum()
    both_wrong = ((a_em == 0) & (b_em == 0)).sum()

    table = [[both_correct, a_only], [b_only, both_wrong]]
    # With zero discordant pairs statsmodels divides by zero; the verdict is withheld anyway
    with np.errstate(divide="ignore", invalid="ignore"):
        result = mcnemar(table, exact=False, correction=True)
    p_value = result.pvalue
    # The chi-square approximation is unreliable with few discordant pairs -> withhold the verdict
    if (a_only + b_only) < min_discordant:
        sig = f"N/A (discordant < {min_discordant})"
    else:
        sig = SIGNIFICANT if p_value < alpha else NOT_SIGNIFICANT

    return {
        "Questions": len(group),
        "Both correct": both_correct,
        f"Only {model_A} correct": a_only,
        f"Only {model_B} correct": b_only,
        "Both wrong": both_wrong,
        "Statistic": result.statistic,
        "p-value": p_value,
        SIG_COL: sig,
    }


def benchmark_mcnemar(cfg, final_df, model_A, model_B) -> pd.DataFrame:
    """Per-benchmark test, plus EXTRA_GROUPS aggregates."""
    def test(group):
        return mcnemar_test(group, model_A, model_B, cfg.ALPHA, cfg.MIN_DISCORDANT)

    rows = [
        {"Benchmark": cfg.BENCHMARK_LABELS.get(name, name), **test(group)}
        for name, group in final_df.groupby("benchmark")
    ]
    for label, spec in cfg.EXTRA_GROUPS.items():
        group = extra_group(final_df, spec)
        if not group.empty:
            rows.append({"Benchmark": label, **test(group)})
    return pd.DataFrame(rows)


def category_mcnemar(cfg, final_df, benchmark_name, model_A, model_B) -> pd.DataFrame:
    """Per-category (sub-task) test within one benchmark."""
    # Filter on the benchmark column so kmmlu does not pick up kmmlu_pro, etc.
    sub_df = final_df[final_df["benchmark"] == benchmark_name]
    return pd.DataFrame([
        {"Category": category, **mcnemar_test(group, model_A, model_B, cfg.ALPHA, cfg.MIN_DISCORDANT)}
        for category, group in sub_df.groupby("category")
    ])


def domain_mcnemar(cfg, final_df, model_A, model_B) -> pd.DataFrame:
    rows = []
    for domain_name in cfg.DOMAINS:
        group = domain_group(cfg, final_df, domain_name)
        if not group.empty:
            rows.append({"Domain": domain_name,
                         **mcnemar_test(group, model_A, model_B, cfg.ALPHA, cfg.MIN_DISCORDANT)})
    return pd.DataFrame(rows)


def print_domain_composition(cfg, final_df: pd.DataFrame):
    print("\nDomain composition (questions):")
    for domain_name in cfg.DOMAINS:
        total_count = 0
        print(f"\n  {domain_name}")
        for key in cfg.DOMAIN_TO_BENCHMARKS.get(domain_name, []):
            group = select_group(final_df, key)
            count = 0 if group is None else group.shape[0]
            total_count += count
            print(f"   ├─ {key:<35} : {count}")
        print(f"   └─ {'total':<35} : {total_count}")


# ── Accuracy ──
def score_row(kind, name, total, a_correct, b_correct, a_acc, b_acc, model_A, model_B):
    return {
        "Type": kind,
        "Name": name,
        "Questions": total,
        correct_col(model_A): a_correct,
        accuracy_col(model_A): a_acc,
        correct_col(model_B): b_correct,
        accuracy_col(model_B): b_acc,
    }


def compute_group_score(group, name, model_A, model_B, kind):
    total = len(group)
    a_correct = group[f"{model_A}_EM"].sum()
    b_correct = group[f"{model_B}_EM"].sum()
    return score_row(kind, name, total, a_correct, b_correct,
                     round(a_correct / total * 100, 2), round(b_correct / total * 100, 2),
                     model_A, model_B)


def compute_ifeval_score(final_df, base_name, model_A, model_B):
    """ifeval-style benchmark: simple mean of the plsa / ilsa / plla / illa accuracies."""
    categories = set(final_df["category"].unique())
    sub_groups = [
        final_df[final_df["category"] == f"{base_name}_{x}"]
        for x in ["illa", "ilsa", "plla", "plsa"]
        if f"{base_name}_{x}" in categories
    ]
    if not sub_groups:
        return None

    avg_a = sum(g[f"{model_A}_EM"].mean() for g in sub_groups) / len(sub_groups)
    avg_b = sum(g[f"{model_B}_EM"].mean() for g in sub_groups) / len(sub_groups)
    total = sum(len(g) for g in sub_groups)
    return score_row("Benchmark", base_name, total, round(avg_a * total), round(avg_b * total),
                     round(avg_a * 100, 2), round(avg_b * 100, 2), model_A, model_B)


def accuracy_table(cfg, final_df, model_A, model_B, significance: dict) -> pd.DataFrame:
    """Accuracy per benchmark / category / domain, with the McNemar verdict attached."""
    score_cache = {}
    rows = []
    benchmark_names = set(final_df["benchmark"].unique())

    def add(row):
        rows.append(row)
        score_cache[row["Name"]] = (row[accuracy_col(model_A)], row[accuracy_col(model_B)])

    # 1. Benchmarks
    for benchmark, group in final_df.groupby("benchmark"):
        if benchmark in IFEVAL_LIKE:
            row = compute_ifeval_score(final_df, benchmark, model_A, model_B)
            if row:
                add(row)
            continue
        label = cfg.BENCHMARK_LABELS.get(benchmark, benchmark)
        add(compute_group_score(group, label, model_A, model_B, "Benchmark"))
        score_cache[benchmark] = score_cache[label]

    # 1-1. EXTRA_GROUPS aggregates
    for label, spec in cfg.EXTRA_GROUPS.items():
        group = extra_group(final_df, spec)
        if not group.empty:
            add(compute_group_score(group, label, model_A, model_B, "Benchmark"))

    # 2. Categories (skipped when identical to a benchmark name)
    for category, group in final_df.groupby("category"):
        if category not in benchmark_names:
            add(compute_group_score(group, category, model_A, model_B, "Category"))

    # 3. Domains: question-weighted (DOMAIN_TO_BENCHMARKS)
    for domain_name in cfg.DOMAIN_TO_BENCHMARKS:
        if domain_name not in cfg.DOMAINS or domain_name in cfg.DOMAIN_AVERAGE:
            continue
        group = domain_group(cfg, final_df, domain_name)
        if not group.empty:
            add(compute_group_score(group, domain_name, model_A, model_B, "Domain"))

    # 4. Domains: simple mean of sub-item scores (DOMAIN_AVERAGE)
    for domain_name, sub_items in cfg.DOMAIN_AVERAGE.items():
        if domain_name not in cfg.DOMAINS:
            continue
        scores_A, scores_B = [], []
        for item in sub_items:
            if item in score_cache:
                a, b = score_cache[item]
            else:
                group = select_group(final_df, item)
                if group is None or group.empty:
                    continue
                a = group[f"{model_A}_EM"].mean() * 100
                b = group[f"{model_B}_EM"].mean() * 100
            scores_A.append(a)
            scores_B.append(b)
        if scores_A:
            add(score_row("Domain", domain_name, "-", "-", "-",
                          round(sum(scores_A) / len(scores_A), 2),
                          round(sum(scores_B) / len(scores_B), 2), model_A, model_B))

    acc_df = pd.DataFrame(rows)
    acc_df[DIFF_COL] = (acc_df[accuracy_col(model_A)] - acc_df[accuracy_col(model_B)]).round(2)
    acc_df[SIG_COL] = [significance.get((t, n), "") for t, n in zip(acc_df["Type"], acc_df["Name"])]
    acc_df["_order"] = acc_df["Type"].map(TYPE_ORDER.index)
    return acc_df.sort_values(by=["_order", "Name"]).drop(columns="_order")


# ── Excel output ──
def _display_width(value) -> int:
    if value is None:
        return 0
    if isinstance(value, float):
        return len(f"{value:.4f}")
    return len(str(value))


def style_sheet(ws, model_A=None, model_B=None, highlight=None):
    """Header style, freeze panes, filter, number formats, column widths, significance highlight.

    highlight: "row" = fill the whole significant row, "cell" = fill only the Significance cell.
    """
    if ws.max_column < 1 or ws.max_row < 1:
        return
    header = [cell.value for cell in ws[1]]

    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.fill = HEADER_FILL
    ws.freeze_panes = "B2"
    ws.auto_filter.ref = ws.dimensions

    accuracy_cols = {accuracy_col(model_A), accuracy_col(model_B)} if model_A else set()
    sig_idx = header.index(SIG_COL) if SIG_COL in header else None

    for col_idx, name in enumerate(header, start=1):
        fmt = "0.00" if name in accuracy_cols else NUMBER_FORMATS.get(name)
        for (cell,) in ws.iter_rows(min_row=2, min_col=col_idx, max_col=col_idx):
            if fmt and isinstance(cell.value, (int, float)):
                cell.number_format = fmt
            if name in accuracy_cols:
                cell.fill = ACCURACY_FILL

    if highlight and sig_idx is not None:
        for row in ws.iter_rows(min_row=2):
            if row[sig_idx].value == SIGNIFICANT:
                for cell in (row if highlight == "row" else [row[sig_idx]]):
                    cell.fill = SIGNIFICANT_FILL

    for col_idx, column_cells in enumerate(ws.columns, start=1):
        width = max(_display_width(c.value) for c in column_cells)
        ws.column_dimensions[get_column_letter(col_idx)].width = min(max(width + 3, 8), 60)


def info_table(cfg, model_A, model_B, compared, skipped) -> pd.DataFrame:
    data_path = cfg.DATA_PATH
    if data_path.is_relative_to(REPO_ROOT):
        data_path = data_path.relative_to(REPO_ROOT)
    items = [
        ("Model A", model_A),
        ("Model B", model_B),
        ("Config", cfg.NAME),
        ("Data path", str(data_path)),
        ("Test", "McNemar's test (chi-square, continuity correction)"),
        ("Alpha", cfg.ALPHA),
        ("Min discordant pairs", cfg.MIN_DISCORDANT),
        ("Generated at", datetime.now().strftime("%Y-%m-%d %H:%M:%S")),
        ("Benchmarks compared", ", ".join(compared)),
        ("Benchmarks skipped", len(skipped)),
    ]
    items += [(f"  skipped: {bench}", reason) for bench, reason in skipped]
    return pd.DataFrame(items, columns=["Item", "Value"])


def write_report(file_path, sheets: dict, model_A, model_B):
    """sheets: {sheet name: (DataFrame, highlight)}. Written in dict order."""
    with pd.ExcelWriter(file_path, engine="openpyxl") as writer:
        for name, (df, _highlight) in sheets.items():
            df.to_excel(writer, sheet_name=name, index=False)
        for name, (_df, highlight) in sheets.items():
            ws = writer.sheets[name]
            if name == SHEET_INFO:
                ws["A1"].font = ws["B1"].font = Font(bold=True)
                ws.column_dimensions["A"].width = 28
                ws.column_dimensions["B"].width = 80
            else:
                style_sheet(ws, model_A, model_B, highlight)


# ── Run ──
def build_paired_df(cfg, model_A: str, model_B: str):
    """Join both models' results per benchmark by question id. Benchmarks that do not pair up are skipped."""
    all_results, skipped = [], []

    for benchmark in cfg.BENCHMARKS:
        try:
            df_A = load_model_result(cfg, benchmark, model_A)
            df_B = load_model_result(cfg, benchmark, model_B)

            if len(df_A) != len(df_B):
                raise ValueError(f"row count mismatch (A={len(df_A)}, B={len(df_B)})")

            for model, df in ((model_A, df_A), (model_B, df_B)):
                if df["id"].duplicated().any():
                    raise ValueError(f"duplicate question ids in {model} results")
                em = df[f"{model}_EM"]
                if not em.isin([0, 1]).all():
                    bad = em[~em.isin([0, 1])].unique()[:3]
                    raise ValueError(f"non-binary scores in {model} results (e.g. {list(bad)}); "
                                     "McNemar's test needs correct/incorrect per question")

            merged = df_A[["id", "category", "benchmark", f"{model_A}_EM"]]
            merged = pd.merge(merged, df_B[["id", f"{model_B}_EM"]], on="id", how="left")

            # Guard against questions that do not line up (different categories/order)
            if merged[f"{model_A}_EM"].isna().any() or merged[f"{model_B}_EM"].isna().any():
                raise ValueError("question ids do not match (different categories or order)")

            all_results.append(merged)
            print(f"[OK]   {benchmark}: {merged.shape[0]} rows")
        except Exception as e:
            print(f"[SKIP] {benchmark}: {e}")
            skipped.append((benchmark, str(e)))

    if not all_results:
        return None, skipped
    final_df = pd.concat(all_results, ignore_index=True)
    return final_df[["id", "benchmark", "category", f"{model_A}_EM", f"{model_B}_EM"]], skipped


def run(cfg, model_A: str, model_B: str, output_dir=RESULT_DIR) -> str | None:
    print("\n" + "=" * 72)
    print(f"McNemar's test  A={model_A}  vs  B={model_B}  ({cfg.DATA_PATH})")
    print("=" * 72)

    final_df, skipped = build_paired_df(cfg, model_A, model_B)
    if final_df is None:
        print("[WARN] no benchmark could be compared")
        return None

    present = set(final_df["benchmark"].unique())
    bench_df = benchmark_mcnemar(cfg, final_df, model_A, model_B)
    category_dfs = {
        bench: category_mcnemar(cfg, final_df, bench, model_A, model_B)
        for bench in cfg.BENCHMARKS
        if bench in cfg.CATEGORY_BENCHMARKS and bench in present
    }
    domain_df = domain_mcnemar(cfg, final_df, model_A, model_B)
    print_domain_composition(cfg, final_df)

    # Verdicts shown next to the accuracies: {(Type, Name): Significance}
    significance = {("Benchmark", r["Benchmark"]): r[SIG_COL] for _, r in bench_df.iterrows()}
    for df in category_dfs.values():
        significance.update({("Category", r["Category"]): r[SIG_COL] for _, r in df.iterrows()})
    significance.update({("Domain", r["Domain"]): r[SIG_COL] for _, r in domain_df.iterrows()})
    acc_df = accuracy_table(cfg, final_df, model_A, model_B, significance)

    compared = [b for b in cfg.BENCHMARKS if b in present]
    sheets = {
        SHEET_ACCURACY: (acc_df, "cell"),
        SHEET_BENCHMARK: (bench_df, "row"),
        SHEET_DOMAIN: (domain_df, "row"),
        **{bench[:31]: (df, "row") for bench, df in category_dfs.items()},
        SHEET_RAW: (final_df, None),
        SHEET_INFO: (info_table(cfg, model_A, model_B, compared, skipped), None),
    }

    output_dir = Path(output_dir)
    output_dir.mkdir(parents=True, exist_ok=True)
    file_path = str(output_dir / build_result_filename(model_A, model_B, cfg.RESULT_SLUG))
    write_report(file_path, sheets, model_A, model_B)
    print(f"\n[DONE] report saved -> {file_path}")

    if skipped:
        print(f"[WARN] {len(skipped)} benchmark(s) skipped: {', '.join(b for b, _ in skipped)}")
    return file_path


def main(argv=None):
    parser = argparse.ArgumentParser(description="McNemar's test between two models' benchmark results")
    parser.add_argument("--config", "-c", required=True, choices=available_configs(),
                        help="config name under src/configs/")
    parser.add_argument("--model-a", "-a", required=True, help="model A folder name")
    parser.add_argument("--model-b", "-b", required=True, help="model B folder name")
    parser.add_argument("--data-dir", help="folder holding the model folders (default: data/{DATA_DIR})")
    parser.add_argument("--output-dir", default=str(RESULT_DIR), help="where to save the report (default: result/)")
    parser.add_argument("--alpha", type=float, help="significance level (default: 0.05)")
    parser.add_argument("--min-discordant", type=int,
                        help="withhold the verdict when discordant pairs are fewer than this (default: 25)")
    args = parser.parse_args(argv)

    cfg = load_config(args.config, args.data_dir)
    if args.alpha is not None:
        cfg.ALPHA = args.alpha
    if args.min_discordant is not None:
        cfg.MIN_DISCORDANT = args.min_discordant

    missing = [m for m in (args.model_a, args.model_b) if not (cfg.DATA_PATH / m).is_dir()]
    if missing:
        models = sorted(p.name for p in cfg.DATA_PATH.iterdir() if p.is_dir()) if cfg.DATA_PATH.is_dir() else []
        sys.exit(f"[ERROR] model folder not found in {cfg.DATA_PATH}: {missing}\navailable: {models}")

    run(cfg, args.model_a, args.model_b, args.output_dir)


if __name__ == "__main__":
    main()
