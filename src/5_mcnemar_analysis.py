import pandas as pd
import numpy as np
import os
from statsmodels.stats.contingency_tables import mcnemar
from openpyxl import load_workbook
from openpyxl.styles import PatternFill

# 엑셀 파일 불러오
folder_path = "./res_merge_data_en/Llama-74B_GPT4o_en_벤치마크"
folder_path_base_name = os.path.basename(folder_path)
excel_files = [f for f in os.listdir(folder_path)]
alpha = 0.05  # 유의수준

for file in excel_files:
    file_path = os.path.join(folder_path, file)
    df = pd.read_excel(file_path, sheet_name="Sheet1")
    model_A_name = df.columns[2]
    model_B_name = df.columns[3]

    # 결과 저장용 리스트
    results = []
    print(df.columns)

    # 전체 맥네마 검정 (전체 문제 기준)
    group1 = df.iloc[:,2]
    group2 = df.iloc[:,3]

    if group1.notna().sum() != group2.notna().sum():
        raise ValueError(f"[에러] 문제 수 다름 → 파일: {file}, group1 (non-NaN): {group1.notna().sum()}, group2 (non-NaN): {group2.notna().sum()}")

    A = ((group1 == 1) & (group2 == 1)).sum()
    B = ((group1 == 1) & (group2 == 0)).sum()
    C = ((group1 == 0) & (group2 == 1)).sum()
    D = ((group1 == 0) & (group2 == 0)).sum()

    table = np.array([[A, B], [C, D]])
    result = mcnemar(table, exact=False, correction=True)

    results.append({
        "A (둘 다 정답)": A,
        f"B ({model_A_name}만 정답)": B,
        f"C ({model_B_name}만 정답)": C,
        "D (둘 다 오답)": D,
        "통계량": result.statistic,
        "p값": result.pvalue,
        "유의성 판단": "-" if (B + C) < 25 else ("유의함" if result.pvalue < alpha else "유의하지 않음")
    })

    # 결과 DataFrame 생성
    result_df = pd.DataFrame(results)

    # 카테고리별 맥네마 검정
    category_results = []
    for category, group in df.groupby('category'):
        group1 = group.iloc[:, 2]
        group2 = group.iloc[:, 3]

        if group1.notna().sum() != group2.notna().sum():
            raise ValueError(f"[에러] 문제 수 다름 (카테고리: {category}) → group1: {group1.notna().sum()}, group2: {group2.notna().sum()}")

        A = ((group1 == 1) & (group2 == 1)).sum()
        B = ((group1 == 1) & (group2 == 0)).sum()
        C = ((group1 == 0) & (group2 == 1)).sum()
        D = ((group1 == 0) & (group2 == 0)).sum()

        table = np.array([[A, B], [C, D]])
        result = mcnemar(table, exact=False, correction=True)

        category_results.append({
            "카테고리": category,
            "A (둘 다 정답)": A,
            f"B ({model_A_name}만 정답)": B,
            f"C ({model_B_name}만 정답)": C,
            "D (둘 다 오답)": D,
            "통계량": result.statistic,
            "p값": result.pvalue,
            "유의성 판단": "-" if (B + C) < 25 else ("유의함" if result.pvalue < alpha else "유의하지 않음")
        })

    category_groups = {
        "kobest_boolq + kobest_sentineg": ["kobest_boolq", "kobest_sentineg"],
        "kobest_copa + kobest_hellaswag + kobest_wic": ["kobest_copa", "kobest_hellaswag", "kobest_wic"],
        "ko_sovereign_culture_folklore + ko_sovereign_society": ["ko_sovereign_culture_folklore", "ko_sovereign_society"],
        "ko_sovereign_영역": ["ko_sovereign_history", "ko_sovereign_law", "ko_sovereign_economics_finance", "ko_sovereign_politics", "ko_sovereign_education", "ko_sovereign_geography"],
    }
    for combined_category_name, combined_categories in category_groups.items():
        combined_df = df[df["category"].isin(combined_categories)]
        if combined_df.empty:
            continue

        group1 = combined_df.iloc[:, 2]
        group2 = combined_df.iloc[:, 3]

        if group1.notna().sum() != group2.notna().sum():
            raise ValueError(f"[에러] 문제 수 다름 (통합 카테고리: {combined_category_name}) → group1: {group1.notna().sum()}, group2: {group2.notna().sum()}")

        A = ((group1 == 1) & (group2 == 1)).sum()
        B = ((group1 == 1) & (group2 == 0)).sum()
        C = ((group1 == 0) & (group2 == 1)).sum()
        D = ((group1 == 0) & (group2 == 0)).sum()

        table = np.array([[A, B], [C, D]])
        result = mcnemar(table, exact=False, correction=True)

        category_results.append({
            "카테고리": combined_category_name,
            "A (둘 다 정답)": A,
            f"B ({model_A_name}만 정답)": B,
            f"C ({model_B_name}만 정답)": C,
            "D (둘 다 오답)": D,
            "통계량": result.statistic,
            "p값": result.pvalue,
            "유의성 판단": "-" if (B + C) < 25 else ("유의함" if result.pvalue < alpha else "유의하지 않음")
        })
    
    category_result_df = pd.DataFrame(category_results)

    # 파일 이름 자동 생성
    base_name, ext = os.path.splitext(file)
    output_path = f"./res_mcnemar/{folder_path_base_name}/{base_name}_mcnemar.xlsx"
    os.makedirs(os.path.dirname(output_path), exist_ok=True)

    with pd.ExcelWriter(output_path, engine="openpyxl") as writer:
        df.to_excel(writer, sheet_name="Sheet1", index=False)
        result_df.to_excel(writer, sheet_name="mcnemar", index=False)
        category_result_df.to_excel(writer, sheet_name="mcnemar_by_category", index=False)

    wb = load_workbook(output_path)
    ws = wb["mcnemar"]
    fill = PatternFill(start_color="FFFF9999", end_color="FFFF9999", fill_type="solid")  # 연한 빨간색

    # '유의성 판단' 열의 위치 찾기 (헤더 기준)
    header = [cell.value for cell in ws[1]]
    judgment_col = header.index("유의성 판단") + 1  # 1-based index

    # 2행부터 체크하면서 "차이 있음"이면 빨간색 배경
    for row in range(2, ws.max_row + 1):
        cell = ws.cell(row=row, column=judgment_col)
        if cell.value == "유의함":
            for col in range(1, ws.max_column + 1):
                ws.cell(row=row, column=col).fill = fill
                print("유의함")

    ws_category = wb["mcnemar_by_category"]
    header_category = [cell.value for cell in ws_category[1]]
    judgment_col_category = header_category.index("유의성 판단") + 1

    for row in range(2, ws_category.max_row + 1):
        cell = ws_category.cell(row=row, column=judgment_col_category)
        if cell.value == "유의함":
            for col in range(1, ws_category.max_column + 1):
                ws_category.cell(row=row, column=col).fill = fill
            print("유의함 (카테고리별 기준)")

    # 저장
    wb.save(output_path)
    print(f"엑셀 저장 완료 (차이 있음: 빨간색 강조): {output_path}")

