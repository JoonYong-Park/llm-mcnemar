import os
import json
import pandas as pd
from statsmodels.stats.contingency_tables import mcnemar
import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import PatternFill
from openpyxl.utils import get_column_letter
import sys
from domain_config_law import (
    BENCHMARKS,
    DOMAIN,
    DOMAIN_NEW,

    DOMAIN_TO_BENCHMARKS,
    DOMAIN_TO_BENCHMARKS_NEW,
    DOMAIN_TO_BENCHMARKS_AVERAGE,
    GROUPED_RAW_CATEGORIES,
    HIERARCHY
)

model_A = "midm-legal"
model_B = "qwen-3-14b"


def load_koifeval_result(model, benchmark_name):
    folder_path = f"./법률지표/{model}/{benchmark_name}/"
    files = [f for f in os.listdir(folder_path) if f.endswith(".jsonl")]
    if not files:
        raise FileNotFoundError(f"[❌ 없음] {folder_path}에 .jsonl 파일이 없습니다.")
    
    filepath = os.path.join(folder_path, files[0])
    rows = []

    with open(filepath, "r", encoding="utf-8") as f:
        for i, line in enumerate(f):
            data = json.loads(line)
            score = data["score"] if "score" in data else data

            rows.append({"id": f"{i}_plsa", "category": "KoIFEval_plsa", f"{model}_EM": score["prompt_level_strict_acc"]})
            for j, val in enumerate(score["inst_level_strict_acc"], 1):
                rows.append({"id": f"{i}-{j}_ilsa", "category": "KoIFEval_ilsa", f"{model}_EM": val})
            rows.append({"id": f"{i}_plla", "category": "KoIFEval_plla", f"{model}_EM": score["prompt_level_loose_acc"]})
            for j, val in enumerate(score["inst_level_loose_acc"], 1):
                rows.append({"id": f"{i}-{j}_illa", "category": "KoIFEval_illa", f"{model}_EM": val})
    
    return pd.DataFrame(rows)


def rename_and_filter_em_column(df: pd.DataFrame, model_name: str) -> pd.DataFrame:
    if 'exact_match' in df.columns:
        df = df.rename(columns={'exact_match': f'{model_name}_EM'})
    elif 'acc' in df.columns:
        df = df.rename(columns={'acc': f'{model_name}_EM'})
    elif 'q_eval' in df.columns:
        df['q_eval'] = df['q_eval'].astype(str).str.lower().str.strip().map({'true': 1, 'false': 0})
        df = df.rename(columns={'q_eval': f'{model_name}_EM'})    
    else:
        raise ValueError(f"[ERROR] {model_name} 결과에 정확도 컬럼이 없습니다.")

    if 'filter' in df.columns:
        filter_values = df['filter'].unique()
        if 'get-answer' in filter_values:
            df = df[df['filter'] == 'get-answer']
        elif 'none' in filter_values:
            df = df[df['filter'] == 'none']
        elif 'custom-extract' in filter_values:
            df = df[df['filter'] == 'custom-extract']
        else:
            pass

    return df

def load_model_result(benchmark_name, model):
    if benchmark_name == "KoIFEval":
        return load_koifeval_result(model, benchmark_name)
    
    df_list = []
    categories = []
    folder_paths = [f"./법률지표/{model}/{benchmark_name}/"]
    for folder_path in folder_paths:
        jsonl_files = [f for f in os.listdir(folder_path) if f.endswith(('.json', '.jsonl', '.xlsx'))]
        for file in jsonl_files:
            file_path = os.path.join(folder_path, file)
            if file.endswith('.json'):
                df = pd.read_json(file_path)
            elif file.endswith('.jsonl'):
                df = pd.read_json(file_path, lines=True, encoding="utf-8-sig")
            elif file.endswith('.xlsx'):
                df = pd.read_excel(file_path)

            base_filename = os.path.splitext(file)[0]
            df['id'] = df.apply(lambda x: file.replace('.jsonl', '').replace('.xlsx', '') + str(x.name), axis=1)
            df['category'] = base_filename
            df = rename_and_filter_em_column(df, model)
            df_list.append(df)
            categories.append(base_filename)
    full_df = pd.concat(df_list, ignore_index=True, sort=False)
    return full_df

def calculate_mcnemar(group: pd.DataFrame, model_A: str, model_B: str, alpha: float = 0.05):
    a_only = ((group[f"{model_A}_EM"] == 1) & (group[f"{model_B}_EM"] == 0)).sum()
    b_only = ((group[f"{model_A}_EM"] == 0) & (group[f"{model_B}_EM"] == 1)).sum()
    both_correct = ((group[f"{model_A}_EM"] == 1) & (group[f"{model_B}_EM"] == 1)).sum()
    both_wrong = ((group[f"{model_A}_EM"] == 0) & (group[f"{model_B}_EM"] == 0)).sum()

    table = [[both_correct, a_only], [b_only, both_wrong]]
    result = mcnemar(table, exact=False, correction=True)
    p_value = result.pvalue
    statistic = result.statistic
    sig = "-" if (a_only + b_only) < 25 else ("유의함" if p_value < alpha else "유의하지 않음")

    return {
        "A (둘 다 정답)": both_correct,
        f"B ({model_A}만 정답)": a_only,
        f"C ({model_B}만 정답)": b_only,
        "D (둘 다 오답)": both_wrong,
        "통계량": statistic,
        "p값": p_value,
        "유의성 판단": sig
    }

def highlight_significant_cells(ws, 판단_col_name="유의성 판단", 색상="FF9999"):

    header = [cell.value for cell in ws[1]]
    if 판단_col_name not in header:
        print(f"[경고] 헤더에 '{판단_col_name}' 없음 → 색상 적용 생략")
        return
    
    judge_col = header.index(판단_col_name) + 1
    red_fill = PatternFill(start_color=색상, end_color=색상, fill_type="solid")

    for row in range(2, ws.max_row + 1):
        cell = ws.cell(row=row, column=judge_col)
        if cell.value == "유의함":
            for col in range(1, ws.max_column + 1):
                ws.cell(row=row, column=col).fill = red_fill

def run_category_mcnemar(df: pd.DataFrame, benchmark_name: str, file_path: str):
    """
    특정 벤치마크 내 카테고리별 유의성 검정을 수행하고 엑셀 저장
    """
    print(f"\n[INFO] {benchmark_name} 카테고리별 유의성 검정 시작")

    sub_df = df[df['benchmark'].str.startswith(f"{benchmark_name}")]

    results = []
    for category, group in sub_df.groupby("category"):
        row = calculate_mcnemar(group, model_A, model_B)
        row = {"카테고리": category, **row}
        results.append(row)

    result_df = pd.DataFrame(results)

    with pd.ExcelWriter(file_path, engine='openpyxl', mode='a', if_sheet_exists='replace') as writer:
        result_df.to_excel(writer, sheet_name=f"{benchmark_name}_by_category", index=False)

    wb = load_workbook(file_path)
    ws = wb[f"{benchmark_name}_by_category"]
    highlight_significant_cells(ws)
    wb.save(file_path)

    print(f"[OK] {benchmark_name} 카테고리별 유의성 검정 완료 → sheet: {benchmark_name}_by_category")

def autofit_column_width(ws):
    """
    워크시트의 각 열 너비를 자동 조정해주는 함수
    """
    for col in ws.columns:
        max_length = 0
        column = col[0].column  # 열 인덱스 (숫자)
        column_letter = get_column_letter(column)  # 열 이름 (예: 'A', 'B', ...)

        for cell in col:
            try:
                if cell.value:
                    max_length = max(max_length, len(str(cell.value)))
            except:
                pass

        adjusted_width = max_length + 2  # 여유 공간 포함
        ws.column_dimensions[column_letter].width = adjusted_width

def compute_accuracy_by_category(df: pd.DataFrame, model_list: list[str], file_path: str):
    print("\n[INFO] category별 모델 정답률 계산 시작")

    results = []

    # 엑셀에 저장
    raw_categories = set(df['raw_category'].unique())
    general_categories = set(df['acc_category'].unique())
    total_categories = raw_categories.union(general_categories - raw_categories)
    category_only_rows = set()

    for category in sorted(total_categories):
    # ✅ ko_comprehensivegen 행은 건너뛰고, 점수는 '한국어생성'에 합산되도록 처리
        if category == "ko_comprehensivegen":
            continue

        row = {"category": category}
        
        if category in df['raw_category'].values:
            sub_df = df[df['raw_category'] == category]

            # ✅ 여기서 'ko_comprehensivegen'도 함께 포함
            if category == "한국어생성":
                extra = df[df['raw_category'] == "ko_comprehensivegen"]
                if not extra.empty:
                    sub_df = pd.concat([sub_df, extra], ignore_index=True)
        
        else:
            sub_df = df[df['acc_category'] == category]
            category_only_rows.add(category)

            # ✅ 여기서도 'ko_comprehensivegen'을 함께 포함
            if category == "한국어생성":
                extra = df[df['raw_category'] == "ko_comprehensivegen"]
                if not extra.empty:
                    sub_df = pd.concat([sub_df, extra], ignore_index=True)

        for model in model_list:
            col_name = f"{model}_EM"
            if col_name in sub_df.columns:
                acc = sub_df[col_name].mean() * 100
                row[model] = round(acc, 2)

        results.append(row)

    
    # for group_name, subcategories in GROUPED_RAW_CATEGORIES.items():
    #     row = {"category": group_name}
    #     sub_df = df[df['raw_category'].isin(subcategories)]
    #     for model in model_list:
    #         col_name = f"{model}_EM"
    #         if col_name in sub_df.columns:
    #             acc = sub_df[col_name].mean() * 100
    #             row[model] = round(acc, 2)
    #     results.append(row)

    acc_df = pd.DataFrame(results)
    with pd.ExcelWriter(file_path, engine='openpyxl', mode='a', if_sheet_exists='replace') as writer:
        acc_df.to_excel(writer, sheet_name="정답률_by_category", index=False)

    # 열 너비 자동 조정 및 저장
    wb = load_workbook(file_path)
    ws = wb["정답률_by_category"]
    blue_fill = PatternFill(start_color="CCECFF", end_color="CCECFF", fill_type="solid")
    for row in ws.iter_rows(min_row=2, max_row=ws.max_row):
        category_cell = row[0]
        if category_cell.value in category_only_rows:
            for cell in row:
                cell.fill = blue_fill

    # 숫자 포맷 적용 (두 번째 열부터)
    for col in range(2, ws.max_column + 1):
        for row in range(2, ws.max_row + 1):
            cell = ws.cell(row=row, column=col)
            cell.number_format = '0.00'

    # 열 너비 자동 조정 함수
    autofit_column_width(ws)

    wb.save(file_path)
    print(f"[OK] category별 모델 정답률 저장 완료 → sheet: 정답률_by_category")

def highlight_accuracy_rows(ws, benchmarks: list[str], domains: list[str], domain_new: list[str]):
    blue_fill = PatternFill(start_color="CCECFF", end_color="CCECFF", fill_type="solid")  # 파란색
    green_fill = PatternFill(start_color="CCFFCC", end_color="CCFFCC", fill_type="solid")  # 초록색

    # 헤더 파악
    header = [cell.value for cell in ws[1]]
    col_map = {name: idx for idx, name in enumerate(header)}

    if not all(key in col_map for key in ["대분류", "중분류", "category"]):
        print("[❌] 대/중분류/카테고리 열이 존재하지 않음")
        return

    for row in ws.iter_rows(min_row=2, max_row=ws.max_row):
        category_cell = row[col_map["category"]]
        category = str(category_cell.value).strip()

        if category in benchmarks or category in domain_new:
            for key in ["대분류", "중분류", "category"]:
                row[col_map[key]].fill = blue_fill

        elif any(category == d or category.startswith(d) for d in domains):
            for key in ["대분류", "중분류", "category"]:
                row[col_map[key]].fill = green_fill

def append_domain_averages_to_accuracy_sheet(file_path: str, sheet_name: str, domain_to_benchmarks_average: dict):
    df = pd.read_excel(file_path, sheet_name=sheet_name)

    domain_rows = []
    for domain, category_list in domain_to_benchmarks_average.items():
        row = {"category": f"{domain}"}
        for col in df.columns[1:]:  # 모델별 정답률 열
            scores = df[df['category'].isin(category_list)][col].dropna().tolist()
            row[col] = round(sum(scores) / len(scores), 2) if scores else None
        domain_rows.append(row)

    df_final = pd.concat([df, pd.DataFrame(domain_rows)], ignore_index=True)

    # 저장
    with pd.ExcelWriter(file_path, engine='openpyxl', mode='a', if_sheet_exists='replace') as writer:
        df_final.to_excel(writer, sheet_name=sheet_name, index=False)

    # 후처리: 열 너비 + 숫자 포맷 + 색상
    wb = load_workbook(file_path)
    ws = wb[sheet_name]
    autofit_column_width(ws)
    for col in range(2, ws.max_column + 1):
        for row in range(2, ws.max_row + 1):
            ws.cell(row=row, column=col).number_format = '0.00'

    # highlight_accuracy_rows(ws, benchmarks=BENCHMARKS, domains=list(domain_to_benchmarks_average.keys()))
    wb.save(file_path)

    print(f"[OK] 도메인별 평균 정답률 추가 완료 → 시트: {sheet_name}")

def add_significance_to_accuracy_sheet(file_path):
    wb = load_workbook(file_path)
    all_sheetnames = wb.sheetnames

    # 유의성 판단 딕셔너리 초기화
    sig_dict = {}

    # 1. 시트 순회 (Sheet1, 정답률_by_category 제외)
    for sheet_name in all_sheetnames:
        if sheet_name in ['Sheet1', '정답률_by_category']:
            continue

        ws = wb[sheet_name]
        header = [cell.value for cell in ws[1]]

        # 어떤 이름 열이 있는지 확인
        if '카테고리' in header:
            key_col = header.index('카테고리') + 1
        elif '벤치마크' in header:
            key_col = header.index('벤치마크') + 1
        elif '지표' in header:
            key_col = header.index('지표') + 1
        else:
            continue  # 해당 시트는 건너뜀

        sig_col = header.index('유의성 판단') + 1 if '유의성 판단' in header else None
        if not sig_col:
            continue  # 판단값이 없으면 무시

        for row in ws.iter_rows(min_row=2, max_row=ws.max_row):
            key = str(row[key_col - 1].value).strip()
            sig_value = row[sig_col - 1].value
            if key and sig_value:
                sig_dict[key] = sig_value  # 매핑 등록

    # 2. 정답률 시트에서 category 열과 매칭해서 추가
    acc_ws = wb["정답률_by_category"]
    acc_header = [cell.value for cell in acc_ws[1]]
    category_col = acc_header.index("category") + 1
    new_col_index = acc_ws.max_column + 1
    acc_ws.cell(row=1, column=new_col_index).value = "유의성 판단"

    for row in range(2, acc_ws.max_row + 1):
        category = str(acc_ws.cell(row=row, column=category_col).value).strip()
        sig = sig_dict.get(category)
        if sig:
            acc_ws.cell(row=row, column=new_col_index).value = sig

    red_fill = PatternFill(start_color="FF9999", end_color="FF9999", fill_type="solid")

    for row in range(2, acc_ws.max_row + 1):
        cell = acc_ws.cell(row=row, column=new_col_index)
        if str(cell.value).strip() == "유의함":
            cell.fill = red_fill

    # 저장
    wb.save(file_path)
    print(f"[OK] '정답률_by_category' 시트에 유의성 판단 열 추가 완료")

def add_hierarchy_columns_to_accuracy_sheet(file_path: str, hierarchy: dict):
    wb = load_workbook(file_path)
    ws = wb["정답률_by_category"]
    data = list(ws.iter_rows(values_only=True))
    header = list(data[0])
    rows = data[1:]

    category_idx = header.index("category")
    significance_idx = header.index("유의성 판단") if "유의성 판단" in header else None

    new_data = []
    for row in rows:
        category = row[category_idx]
        big_class = mid_class = category  # 기본값

        found = False
        for big, mids in hierarchy.items():
            for mid, cats in mids.items():
                if category in cats:
                    big_class = big
                    mid_class = mid
                    found = True
                    break
            if found:
                break

        new_data.append((big_class, mid_class, *row))

    # 대분류 기준 정렬
    new_data.sort(key=lambda x: (x[0], x[1]))

    # 새 시트 덮어쓰기
    new_ws = wb.create_sheet("정답률_by_category_tmp")
    new_header = ["대분류", "중분류"] + header
    new_ws.append(new_header)

    red_fill = PatternFill(start_color="FF9999", end_color="FF9999", fill_type="solid")

    for row in new_data:
        new_row = []
        for i, val in enumerate(row):
            new_row.append(val)
        new_ws.append(new_row)

    # 스타일 적용
    for row in new_ws.iter_rows(min_row=2, max_row=new_ws.max_row):
        for cell in row:
            if isinstance(cell.value, float):
                cell.number_format = '0.00'
        
        if significance_idx is not None:
            sig_value = row[significance_idx + 2].value  # 기존 열이 2칸 밀렸음
            if sig_value == "유의함":
                row[significance_idx + 2].fill = red_fill

    # 기존 시트 삭제 후 새 시트 이름 변경
    wb.remove(wb["정답률_by_category"])
    new_ws.title = "정답률_by_category"
    highlight_accuracy_rows(
    new_ws, 
    benchmarks=BENCHMARKS, 
    domain_new=DOMAIN_NEW, 
    domains=DOMAIN
)
    wb.save(file_path)
    print(f"[OK] 대/중분류 열 추가 및 스타일 적용 완료 → 시트: 정답률_by_category")

all_results = []

for benchmark in BENCHMARKS:
    try:
        df_A = load_model_result(benchmark, model_A)
        df_B = load_model_result(benchmark, model_B)

        if len(df_A) != len(df_B):
            raise ValueError(f"[❌ 행 개수 불일치] {benchmark}: A={len(df_A)}행, B={len(df_B)}행")

        raw_df = df_A[['id', 'category']].copy()
        raw_df['benchmark'] = benchmark

        merged = pd.merge(raw_df, df_A[['id', f"{model_A}_EM"]], on="id", how="left")
        merged = pd.merge(merged, df_B[['id', f"{model_B}_EM"]], on="id", how="left")

        all_results.append(merged)
        print(f"[OK] {benchmark}: {merged.shape[0]} rows")
    except Exception as e:
        print(f"[ERROR] {benchmark}: {e}")
        sys.exit(1)


if all_results:
    final_df = pd.concat(all_results, ignore_index=True)
    final_df = final_df[['id', 'benchmark', 'category', f'{model_A}_EM', f'{model_B}_EM']]
    os.makedirs("./result", exist_ok=True)
    file_path = f"./result/{model_A}_{model_B}_법률모델_유의성검증_통합결과.xlsx"
    final_df.to_excel(file_path, index=False)
    print(f"[완료] 저장됨 → {file_path}")

    ### ✅ 여기부터 유의성검정 코드 작성 시작 ###
    mcnemar_results = []

    for benchmark_name, group in final_df.groupby("benchmark"):
        row = calculate_mcnemar(group, model_A, model_B)
        row = {"벤치마크": benchmark_name, **row}

        mcnemar_results.append(row)

    mcnemar_df = pd.DataFrame(mcnemar_results)

    with pd.ExcelWriter(file_path, engine='openpyxl', mode='a', if_sheet_exists='replace') as writer:
        mcnemar_df.to_excel(writer, sheet_name="benchmark_mcnemar", index=False)

    print(f"[OK] 유의성검정 결과 저장 완료 → sheet: benchmark_mcnemar")

    wb = load_workbook(file_path)
    ws = wb["benchmark_mcnemar"]

    highlight_significant_cells(ws)
    autofit_column_width(ws)

    wb.save(file_path)
    print(f"[OK] '유의함' 셀 빨간색 색칠 완료 → {file_path}")

else:
    print("[경고] 결과 없음")

for bench in BENCHMARKS:
    if bench in ['법률지식_QA','법률추론','한국어이해','한국어전문지식','KoIFEval']:
        run_category_mcnemar(final_df, bench, file_path)

final_df['raw_category'] = final_df['category']

def extract_benchmark(category):

    if category.startswith(('K-Honorifics-hard', 'K-Implicature-hard', 'K-Lexicon-hard')):
        return 'kpragmatics_hard'
    elif category.startswith(('k_honorifics', 'k_implicature', 'k_lexicon')):
        return 'kpragmatics'
    elif category.startswith(('KoIFEval')):
        return 'KoIFEval'
    else:
        return category
    
def acc_benchmark(category):
    if category.startswith(('K-Honorifics-hard', 'K-Implicature-hard', 'K-Lexicon-hard')):
        return 'kpragmatics_hard'
    # elif category.startswith(('k_honorifics', 'k_implicature', 'k_lexicon')):
    #     return 'kpragmatics'
    elif category.startswith(('ko_comprehensivegen')):
        return '한국어생성'
    # elif category.startswith(('ko_sovereign_')):
    #     return 'ko_sovereign'
    else:
        return category


final_df['acc_category'] = final_df['category'].apply(acc_benchmark)
final_df['category'] = final_df['category'].apply(extract_benchmark)

# domain_results = []

# for domain_name in DOMAIN:
#     categories = DOMAIN_TO_BENCHMARKS.get(domain_name, [])
#     group = final_df[final_df['category'].isin(categories)]

#     if group.empty:
#         continue

#     row = calculate_mcnemar(group, model_A, model_B)
#     row = {"지표": domain_name, **row}
#     domain_results.append(row)

# # ✅ Sheet3 저장
# domain_df = pd.DataFrame(domain_results)

# with pd.ExcelWriter(file_path, engine='openpyxl', mode='a', if_sheet_exists='replace') as writer:
#     domain_df.to_excel(writer, sheet_name="지표_mcnemar", index=False)

# # ✅ "유의함" 빨간 배경 표시
# wb = load_workbook(file_path)
# ws = wb["지표_mcnemar"]
# highlight_significant_cells(ws)
# autofit_column_width(ws)

# wb.save(file_path)
# print(f"[OK] 지표별(category 기준) 유의성 검정 완료 → sheet: 지표_mcnemar")

# print("\n✅ 지표별 카테고리 구성 및 문제 수:")
# for domain_name in DOMAIN:
#     categories = DOMAIN_TO_BENCHMARKS.get(domain_name, [])
#     total_count = 0
#     print(f"\n📌 지표: {domain_name}")
#     for cat in categories:
#         count = final_df[final_df['category'] == cat].shape[0]
#         total_count += count
#         print(f"   ├─ {cat:<35} : {count}개")
#     print(f"   └─ 총합{'':<33} : {total_count}개")

# try:
#     new_domain_results = []

#     for domain_name in DOMAIN_NEW:
#         categories = DOMAIN_TO_BENCHMARKS_NEW[domain_name]
#         group = final_df[final_df['raw_category'].isin(categories)]

#         if group.empty:
#             continue

#         row = calculate_mcnemar(group, model_A, model_B)
#         row = {"KMMLU_지표": domain_name, **row}
#         new_domain_results.append(row)

#     # ✅ 결과를 기존 시트에 이어붙이기
#     new_df = pd.DataFrame(new_domain_results)
#     wb = load_workbook(file_path)
#     ws = wb["지표_mcnemar"]
#     from openpyxl.utils.dataframe import dataframe_to_rows
#     for row in dataframe_to_rows(new_df, index=False, header=True):
#         ws.append(row)

#     highlight_significant_cells(ws)
#     autofit_column_width(ws)

#     wb.save(file_path)
#     print(f"[OK] DOMAIN_TO_BENCHMARKS_NEW 결과 → sheet: 지표_mcnemar 이어붙이기 완료")

# except NameError as e:
#     print(f"[SKIP] DOMAIN_TO_BENCHMARKS_NEW 분석은 건너뜀: {e}")


model_list = [model_A, model_B]  
compute_accuracy_by_category(final_df, model_list, file_path)
append_domain_averages_to_accuracy_sheet(file_path, "정답률_by_category", DOMAIN_TO_BENCHMARKS_AVERAGE)
add_significance_to_accuracy_sheet(file_path)
add_hierarchy_columns_to_accuracy_sheet(file_path, HIERARCHY)
wb = load_workbook(file_path)


desired_order = ["Sheet1", "benchmark_mcnemar", "정답률_by_category"]
ordered_sheets = [ws for name in desired_order for ws in wb.worksheets if ws.title == name]
remaining_sheets = [ws for ws in wb.worksheets if ws.title not in desired_order]

wb._sheets = ordered_sheets + remaining_sheets

wb.save(file_path)
print("[OK] 시트 순서 정리 완료")